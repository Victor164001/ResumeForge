import io

from docx import Document

from app.models.schemas import ContactField, Entry, Heading, Resume, Section
from app.services.model_builder import build_model
from app.services.docx_export import render_docx


def _resume(template="classic", page_size="letter", style="bullets") -> Resume:
    return Resume(
        templateId=template, pageSize=page_size,
        heading=Heading(name="Alex Rivera", role="Senior Product Manager", align="auto",
                         contacts=[ContactField(label="Email", value="alex@example.com")]),
        sections=[
            Section(type="summary", title="Summary", align="auto", text="Product manager, 6 years."),
            Section(type="skills", title="Skills", align="auto", items=["SQL", "Figma"]),
            Section(type="experience", title="Experience", align="auto", style=style, entries=[
                Entry(heading="Senior PM", org="Acme Inc", location="Remote",
                      startMonth="Jan", startYear="2023", isPresent=True,
                      bullets=["Led a platform rewrite", "Grew activation 28%"]),
            ]),
        ],
    )


def _full_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def test_produces_a_valid_docx_with_real_text():
    docx_bytes = render_docx(build_model(_resume()))
    assert docx_bytes[:2] == b"PK"
    text = _full_text(docx_bytes)
    assert "Alex Rivera" in text
    assert "SUMMARY" in text
    assert "Present" in text


def test_works_for_every_template_and_page_size():
    for template in ("classic", "harvard", "modern"):
        for page_size in ("letter", "a4"):
            docx_bytes = render_docx(build_model(_resume(template=template, page_size=page_size)))
            assert docx_bytes[:2] == b"PK", f"{template}/{page_size} failed"


def test_bullets_use_the_list_bullet_style():
    docx_bytes = render_docx(build_model(_resume(style="bullets")))
    doc = Document(io.BytesIO(docx_bytes))
    bullet_paragraphs = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert len(bullet_paragraphs) == 2
    assert bullet_paragraphs[0].text == "Led a platform rewrite"


def test_paragraph_mode_joins_bullets_into_one_flowing_paragraph():
    docx_bytes = render_docx(build_model(_resume(style="paragraph")))
    doc = Document(io.BytesIO(docx_bytes))
    bullet_paragraphs = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert bullet_paragraphs == []
    text = _full_text(docx_bytes)
    assert "Led a platform rewrite. Grew activation 28%." in text


def test_modern_template_colors_the_name_with_the_accent():
    m = build_model(_resume(template="modern"))
    m["accent"] = "#123456"
    docx_bytes = render_docx(m)
    doc = Document(io.BytesIO(docx_bytes))
    name_run = doc.paragraphs[0].runs[0]
    assert str(name_run.font.color.rgb) == "123456"


def test_empty_resume_still_produces_a_valid_docx():
    docx_bytes = render_docx(build_model(Resume(heading=Heading(), sections=[])))
    assert docx_bytes[:2] == b"PK"
