"""
Server-side DOCX export.

Ports the frontend's exportDocx() (docx.js, ui_part2.js): same page setup,
same per-template fonts and accent color, same tab-stop split-row layout
for title/date and org/location lines, same bullets/paragraph modes.
"""
from __future__ import annotations

import io
from typing import Any, Dict

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from app.services.dates import join_as_paragraph

ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _font_for(template: str) -> str:
    if template == "harvard":
        return "Georgia"
    if template == "modern":
        return "Calibri"
    return "Arial"


def _rgb(hex_color: str) -> RGBColor:
    h = (hex_color or "#B23A2F").lstrip("#")
    if len(h) != 6:
        h = "B23A2F"
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _add_bottom_border(paragraph, color_hex: str):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color_hex.lstrip("#"))
    borders.append(bottom)
    p_pr.append(borders)


def _set_run(run, text: str, font: str, size: float, bold=False, italic=False, color: RGBColor | None = None):
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _add_right_tab(paragraph, page_width_pt: float, margin_pt: float):
    tab_stops = paragraph.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Pt(page_width_pt - margin_pt * 2), WD_TAB_ALIGNMENT.RIGHT)


def render_docx(model: Dict[str, Any]) -> bytes:
    template = model["template"]
    is_modern = template == "modern"
    font = _font_for(template)
    accent = _rgb(model.get("accent", "#B23A2F"))
    ink = RGBColor(0x11, 0x11, 0x11)

    page_is_a4 = model.get("page_size") == "a4"
    page_w = Cm(21.0) if page_is_a4 else Cm(21.59)
    page_h = Cm(29.7) if page_is_a4 else Cm(27.94)
    margin = Cm(1.5)

    doc = Document()
    section = doc.sections[0]
    section.page_width = page_w
    section.page_height = page_h
    section.top_margin = section.bottom_margin = margin
    section.left_margin = section.right_margin = margin
    usable_width_pt = page_w.pt - margin.pt * 2

    header = model["header"]
    h_align = ALIGN_MAP.get(header["align"], WD_ALIGN_PARAGRAPH.LEFT)

    p = doc.add_paragraph()
    p.alignment = h_align
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    _set_run(run, header.get("name") or "Your Name", font, 20 if is_modern else 19,
             bold=True, color=accent if is_modern else ink)

    if header.get("role"):
        p = doc.add_paragraph()
        p.alignment = h_align
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run()
        _set_run(run, header["role"], font, 11.5)

    contacts = "   |   ".join(header.get("contacts", []))
    if contacts:
        p = doc.add_paragraph()
        p.alignment = h_align
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run()
        _set_run(run, contacts, font, 9.5, color=RGBColor(0x33, 0x33, 0x33))

    for s in model["sections"]:
        title_align = ALIGN_MAP.get(s["title_align"], WD_ALIGN_PARAGRAPH.LEFT)
        body_align = ALIGN_MAP.get(s["body_align"], WD_ALIGN_PARAGRAPH.LEFT)

        p = doc.add_paragraph()
        p.alignment = title_align
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(4)
        _add_bottom_border(p, model.get("accent", "#B23A2F") if is_modern else "999999")
        run = p.add_run()
        _set_run(run, (s["title"] or "").upper(), font, 10.5, bold=True, color=accent if is_modern else ink)

        if s["type"] == "summary":
            p = doc.add_paragraph()
            p.alignment = body_align
            p.paragraph_format.space_after = Pt(6)
            _set_run(p.add_run(), s.get("text", ""), font, 10)
        elif s["type"] == "skills":
            p = doc.add_paragraph()
            p.alignment = body_align
            p.paragraph_format.space_after = Pt(6)
            _set_run(p.add_run(), "   \u2022   ".join(s.get("items", [])), font, 10)
        else:
            for e in s.get("entries", []):
                if s["split"]:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(1)
                    _add_right_tab(p, usable_width_pt, 0)
                    _set_run(p.add_run(), e["heading"], font, 10.5, bold=True)
                    if e.get("dates"):
                        p.add_run("\t")
                        _set_run(p.add_run(), e["dates"], font, 9.5)
                    if e.get("org") or e.get("location"):
                        p2 = doc.add_paragraph()
                        p2.paragraph_format.space_after = Pt(2)
                        _add_right_tab(p2, usable_width_pt, 0)
                        _set_run(p2.add_run(), e.get("org", ""), font, 9.8, italic=True)
                        if e.get("location"):
                            p2.add_run("\t")
                            _set_run(p2.add_run(), e["location"], font, 9.8, italic=True)
                else:
                    p = doc.add_paragraph()
                    p.alignment = body_align
                    p.paragraph_format.space_after = Pt(1)
                    line1 = "  |  ".join(filter(None, [e["heading"], e.get("dates", "")]))
                    _set_run(p.add_run(), line1, font, 10.5, bold=True)
                    line2 = "  |  ".join(filter(None, [e.get("org", ""), e.get("location", "")]))
                    if line2:
                        p2 = doc.add_paragraph()
                        p2.alignment = body_align
                        p2.paragraph_format.space_after = Pt(2)
                        _set_run(p2.add_run(), line2, font, 9.8, italic=True)

                if e.get("bullets"):
                    if s.get("style") == "paragraph":
                        text = join_as_paragraph(e["bullets"])
                        if text:
                            p = doc.add_paragraph()
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if s["split"] else body_align
                            p.paragraph_format.space_after = Pt(1)
                            _set_run(p.add_run(), text, font, 9.5)
                    else:
                        for b in e["bullets"]:
                            p = doc.add_paragraph(style="List Bullet")
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if s["split"] else body_align
                            p.paragraph_format.space_after = Pt(1)
                            _set_run(p.add_run(), b, font, 9.5)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
